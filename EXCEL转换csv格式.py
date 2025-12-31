import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.styles.numbers import is_date_format
from datetime import datetime
from tkinter import Tk, filedialog, messagebox

def normalize_headers(headers):
    """
    处理重复 / 空列名，确保 JSON 可用
    """
    seen = {}
    new_headers = []

    for i, h in enumerate(headers):
        h = str(h).strip()
        if not h:
            h = f"列{i+1}"

        if h in seen:
            seen[h] += 1
            h = f"{h}_{seen[h]}"
        else:
            seen[h] = 1

        new_headers.append(h)

    return new_headers

def format_cell(cell):
    if cell.value is None:
        return ""
    fmt = cell.number_format
    val = cell.value
    try:
        if is_date_format(fmt) and isinstance(val, datetime):
            return val.strftime("%Y-%m-%d")
        elif "%" in fmt:
            return f"{val * 100:.2f}%"
        elif isinstance(val, float):
            if "." in fmt:
                decimals = fmt[::-1].find(".")
                return f"{val:.{decimals}f}"
            else:
                return str(round(val, 2))
        else:
            return str(val)
    except:
        return str(val)

def unmerge_and_fill_preserve_format(sheet):
    merged_ranges = list(sheet.merged_cells.ranges)
    for merge_range in merged_ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merge_range))
        value = sheet.cell(row=min_row, column=min_col).value
        sheet.unmerge_cells(str(merge_range))
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                if sheet.cell(row=row, column=col).value is None:
                    sheet.cell(row=row, column=col).value = value

    formatted_data = []
    for row in sheet.iter_rows():
        formatted_row = [format_cell(cell) for cell in row]
        formatted_data.append(formatted_row)
    return formatted_data

def process_excel_file(file_path):
    wb = load_workbook(file_path, data_only=True)
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    output_dir = os.path.dirname(file_path)
    exported_files = []
    md_lines = []

    for i, sheetname in enumerate(wb.sheetnames):
        sheet = wb[sheetname]
        formatted_data = unmerge_and_fill_preserve_format(sheet)
        df = pd.DataFrame(formatted_data)
        output_file = os.path.join(output_dir, f"{base_name}_Sheet{i+1}.csv")
        df.to_csv(output_file, index=False, header=False, encoding='utf-8-sig')
        exported_files.append(output_file)

        safe_sheetname = "".join(c if c not in r'\\/:"*?<>|' else "_" for c in sheetname)
        if formatted_data:
            raw_headers = [str(h) for h in formatted_data[0]]
            headers = normalize_headers(raw_headers)
            json_df = pd.DataFrame(formatted_data[1:], columns=headers)
        else:
            headers = []
            json_df = pd.DataFrame()
        json_file = os.path.join(output_dir, f"{base_name}_{safe_sheetname}.json")
        json_df.to_json(json_file, orient="records", force_ascii=False)
        exported_files.append(json_file)

        md_lines.append(f"### Sheet{i+1}: {sheetname}")
        if formatted_data:
            md_lines.append("| " + " | ".join(headers) + " |")
            md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
            for row in formatted_data[1:]:
                md_lines.append("| " + " | ".join(str(cell) for cell in row) + " |")
        md_lines.append("")

    # 写入markdown文件
    md_file = os.path.join(output_dir, f"{base_name}.md")
    with open(md_file, "w", encoding="utf-8") as f:
        f.write("\n".join(md_lines))
    exported_files.append(md_file)

    return exported_files

def choose_file_or_folder():
    root = Tk()
    root.withdraw()

    answer = messagebox.askyesno("选择模式", "是否批量处理整个文件夹？\n是：处理文件夹中所有Excel文件\n否：选择一个文件")

    exported_files = []
    if answer:
        folder = filedialog.askdirectory(title="请选择文件夹")
        if folder:
            for file in os.listdir(folder):
                if file.endswith(".xlsx") or file.endswith(".xlsm"):
                    file_path = os.path.join(folder, file)
                    exported_files += process_excel_file(file_path)
    else:
        file_path = filedialog.askopenfilename(title="请选择Excel文件", filetypes=[("Excel files", "*.xlsx *.xlsm")])
        if file_path:
            exported_files += process_excel_file(file_path)

    if exported_files:
        messagebox.showinfo("处理完成", f"已导出文件：\n" + "\n".join(exported_files))
    else:
        messagebox.showinfo("没有导出文件", "未导出任何文件")

if __name__ == "__main__":
    choose_file_or_folder()
