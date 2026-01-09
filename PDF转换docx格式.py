import os
import tkinter as tk
from tkinter import filedialog, messagebox

import fitz
from docx import Document


def select_path():
    root = tk.Tk()
    root.withdraw()
    choice = messagebox.askyesno("选择方式", "是否处理整个文件夹？\n是 = 文件夹，否 = 单个文件")
    if choice:
        folder = filedialog.askdirectory(title="选择包含 PDF 文件的文件夹")
        return "folder", folder if folder else None
    file_path = filedialog.askopenfilename(title="选择 PDF 文件", filetypes=[("PDF 文件", "*.pdf")])
    return "file", file_path if file_path else None


def extract_text_from_pdf(pdf_path):
    pdf_document = fitz.open(pdf_path)
    pages_text = []
    for page in pdf_document:
        page_text = page.get_text("text").strip()
        pages_text.append(page_text)
    pdf_document.close()
    return pages_text


def convert_pdf_to_docx(pdf_path):
    pages_text = extract_text_from_pdf(pdf_path)
    doc = Document()
    for page_index, page_text in enumerate(pages_text):
        if page_text:
            for line in page_text.splitlines():
                clean_line = line.strip()
                if clean_line:
                    doc.add_paragraph(clean_line)
        else:
            doc.add_paragraph("")
        if page_index < len(pages_text) - 1:
            doc.add_page_break()

    docx_path = os.path.splitext(pdf_path)[0] + ".docx"
    doc.save(docx_path)

    print(f"✅ 已处理：{os.path.basename(pdf_path)}")
    print(f"   ➤ DOCX 输出：{docx_path}")


def main():
    mode, path = select_path()
    if not path:
        print("❌ 未选择路径。")
        return

    if mode == "file":
        convert_pdf_to_docx(path)
    elif mode == "folder":
        files = [f for f in os.listdir(path) if f.lower().endswith(".pdf")]
        if not files:
            print("❌ 文件夹中无 .pdf 文件。")
            return
        for file in files:
            full_path = os.path.join(path, file)
            convert_pdf_to_docx(full_path)


if __name__ == "__main__":
    main()
