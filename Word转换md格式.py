import os
import csv
import tkinter as tk
from tkinter import filedialog, messagebox
from docx import Document as DocxDocument
from docx2python import docx2python

def select_path():
    root = tk.Tk()
    root.withdraw()
    choice = messagebox.askyesno("选择方式", "是否处理整个文件夹？\n是 = 文件夹，否 = 单个文件")
    if choice:
        folder = filedialog.askdirectory(title="选择包含 Word 文件的文件夹")
        return 'folder', folder if folder else None
    else:
        file_path = filedialog.askopenfilename(title="选择 Word 文件", filetypes=[("Word 文件", "*.docx")])
        return 'file', file_path if file_path else None

def flatten_docx2_text(nested_list):
    flat = []
    for item in nested_list:
        if isinstance(item, list):
            flat.extend(flatten_docx2_text(item))
        elif isinstance(item, str) and item.strip():
            flat.append(item.strip())
    return flat

def format_markdown_table(table_rows, title=None):
    """格式化为 Markdown 表格"""
    if not table_rows or len(table_rows) < 2:
        return ""
    header = table_rows[0]
    data_rows = table_rows[1:]

    md_lines = []
    if title:
        md_lines.append(f"### {title}\n\n")
    md_lines.append("| " + " | ".join(header) + " |")
    md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    for row in data_rows:
        md_lines.append("| " + " | ".join(row) + " |")
    md_lines.append("\n")
    return "\n".join(md_lines)

def extract_from_docx(docx_path):
    md_lines = []

    # 提取正文段落
    with docx2python(docx_path) as doc:
        paragraphs = flatten_docx2_text(doc.body)
        if paragraphs:
            md_lines.append("## 正文内容\n\n")
            for para in paragraphs:
                md_lines.append(para + "\n\n")

    # 提取表格内容
    docx = DocxDocument(docx_path)
    for idx, table in enumerate(docx.tables):
        rows = []
        max_cols = 0
        for row in table.rows:
            cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            max_cols = max(max_cols, len(cells))
            rows.append(cells)

        # 补齐每行列数
        for row in rows:
            while len(row) < max_cols:
                row.append("")

        # 添加 Markdown 表格
        table_title = f"表格 {idx + 1}"
        md_table = format_markdown_table(rows, table_title)
        md_lines.append(md_table)

        # 输出 CSV 文件
        base_name = os.path.splitext(os.path.basename(docx_path))[0]
        csv_filename = f"{base_name}_table{idx+1}.csv"
        csv_path = os.path.join(os.path.dirname(docx_path), csv_filename)
        with open(csv_path, "w", newline='', encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            for row in rows:
                writer.writerow(row)

    # 保存 Markdown 文件
    md_path = os.path.splitext(docx_path)[0] + ".md"
    with open(md_path, "w", encoding="utf-8") as md_file:
        md_file.writelines(md_lines)

    print(f"✅ 已处理：{os.path.basename(docx_path)}")
    print(f"   ➤ Markdown 输出：{md_path}")
    print(f"   ➤ 表格数：{len(docx.tables)}")

def main():
    mode, path = select_path()
    if not path:
        print("❌ 未选择路径。")
        return

    if mode == 'file':
        extract_from_docx(path)
    elif mode == 'folder':
        files = [f for f in os.listdir(path) if f.lower().endswith('.docx')]
        if not files:
            print("❌ 文件夹中无 .docx 文件。")
            return
        for file in files:
            full_path = os.path.join(path, file)
            extract_from_docx(full_path)

if __name__ == "__main__":
    main()
