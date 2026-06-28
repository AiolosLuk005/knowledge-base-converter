import os
import csv
import zipfile
import tkinter as tk
from tkinter import filedialog, messagebox
from docx import Document as DocxDocument
from docx2python import docx2python

def is_valid_docx(file_path):
    """验证文件是否为有效的 DOCX 文件"""
    try:
        # 检查文件是否存在
        if not os.path.exists(file_path):
            return False
        # 检查文件大小（空文件或过小的文件不是有效的 DOCX）
        if os.path.getsize(file_path) < 100:
            return False
        # 尝试作为 ZIP 文件打开
        with zipfile.ZipFile(file_path, 'r') as zip_ref:
            # 检查是否包含 DOCX 必需的文件
            namelist = zip_ref.namelist()
            return '[Content_Types].xml' in namelist and 'word/document.xml' in namelist
    except (zipfile.BadZipFile, OSError, PermissionError):
        return False

def select_path():
    root = tk.Tk()
    root.withdraw()
    choice = messagebox.askyesno("选择方式", "是否处理整个文件夹？\n是 = 文件夹，否 = 单个文件")
    if choice:
        folder = filedialog.askdirectory(title="选择包含 Word 文件的文件夹")
        return 'folder', folder if folder else None
    else:
        file_path = filedialog.askopenfilename(title="选择 Word 文件", filetypes=[("Word 文件", "*.docx")])
        return 'file', file_path if file_path else None

def flatten_docx2_text(nested_list):
    flat = []
    for item in nested_list:
        if isinstance(item, list):
            flat.extend(flatten_docx2_text(item))
        elif isinstance(item, str) and item.strip():
            flat.append(item.strip())
    return flat

def format_markdown_table(table_rows, title=None):
    """格式化为 Markdown 表格"""
    if not table_rows or len(table_rows) < 2:
        return ""
    header = table_rows[0]
    data_rows = table_rows[1:]

    md_lines = []
    if title:
        md_lines.append(f"### {title}\n\n")
    md_lines.append("| " + " | ".join(header) + " |")
    md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    for row in data_rows:
        md_lines.append("| " + " | ".join(row) + " |")
    md_lines.append("\n")
    return "\n".join(md_lines)

def extract_from_docx(docx_path):
    """提取 DOCX 文件内容并转换为 Markdown 和 CSV"""
    try:
        md_lines = []

        # 提取正文段落
        try:
            with docx2python(docx_path) as doc:
                paragraphs = flatten_docx2_text(doc.body)
                if paragraphs:
                    md_lines.append("## 正文内容\n\n")
                    for para in paragraphs:
                        md_lines.append(para + "\n\n")
        except Exception as e:
            print(f"⚠️  提取正文时出错：{e}")
            md_lines.append("## 正文内容\n\n")
            md_lines.append("（正文提取失败）\n\n")

        # 提取表格内容
        try:
            docx = DocxDocument(docx_path)
            for idx, table in enumerate(docx.tables):
                rows = []
                max_cols = 0
                for row in table.rows:
                    cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    max_cols = max(max_cols, len(cells))
                    rows.append(cells)

                # 补齐每行列数
                for row in rows:
                    while len(row) < max_cols:
                        row.append("")

                # 添加 Markdown 表格
                table_title = f"表格 {idx + 1}"
                md_table = format_markdown_table(rows, table_title)
                md_lines.append(md_table)

                # 输出 CSV 文件
                base_name = os.path.splitext(os.path.basename(docx_path))[0]
                csv_filename = f"{base_name}_table{idx+1}.csv"
                csv_path = os.path.join(os.path.dirname(docx_path), csv_filename)
                with open(csv_path, "w", newline='', encoding="utf-8-sig") as f:
                    writer = csv.writer(f)
                    for row in rows:
                        writer.writerow(row)

            # 保存 Markdown 文件
            md_path = os.path.splitext(docx_path)[0] + ".md"
            with open(md_path, "w", encoding="utf-8") as md_file:
                md_file.writelines(md_lines)

            print(f"✅ 已处理：{os.path.basename(docx_path)}")
            print(f"   ➤ Markdown 输出：{md_path}")
            print(f"   ➤ 表格数：{len(docx.tables)}")
        except Exception as e:
            print(f"⚠️  提取表格时出错：{e}")
            # 仍然尝试保存已提取的正文内容
            md_path = os.path.splitext(docx_path)[0] + ".md"
            with open(md_path, "w", encoding="utf-8") as md_file:
                md_file.writelines(md_lines)
            print(f"✅ 已保存部分内容：{os.path.basename(docx_path)}")

    except Exception as e:
        print(f"❌ 处理文件失败：{os.path.basename(docx_path)}")
        print(f"   错误详情：{str(e)}")
        return False
    
    return True


def main():
    mode, path = select_path()
    if not path:
        print("❌ 未选择路径。")
        return

    if mode == 'file':
        # 处理单个文件
        if not is_valid_docx(path):
            print(f"❌ 文件不是有效的 DOCX 格式：{os.path.basename(path)}")
            print("   可能原因：")
            print("   - 文件是旧版 .doc 格式（请用 Word 另存为 .docx）")
            print("   - 文件已损坏")
            print("   - 文件扩展名被错误修改")
            return
        extract_from_docx(path)
    elif mode == 'folder':
        # 获取文件夹中的所有 .docx 文件
        all_files = [f for f in os.listdir(path) if f.lower().endswith('.docx')]
        
        # 过滤掉 Word 临时文件（~$ 开头）
        files = [f for f in all_files if not os.path.basename(f).startswith('~$')]
        
        if not files:
            if all_files:
                print("⚠️  文件夹中只有 Word 临时文件（~$ 开头），无有效的 .docx 文件。")
            else:
                print("❌ 文件夹中无 .docx 文件。")
            return
        
        print(f"📁 找到 {len(files)} 个 .docx 文件")
        if len(all_files) > len(files):
            print(f"   (已过滤 {len(all_files) - len(files)} 个临时文件)")
        print("-" * 50)
        
        success_count = 0
        failed_count = 0
        skipped_count = 0
        
        for file in files:
            full_path = os.path.join(path, file)
            
            # 验证文件是否为有效的 DOCX
            if not is_valid_docx(full_path):
                print(f"⚠️  跳过无效文件：{file}")
                print(f"   (文件可能是 .doc 格式或已损坏)")
                skipped_count += 1
                continue
            
            # 处理文件
            if extract_from_docx(full_path):
                success_count += 1
            else:
                failed_count += 1
            print()  # 空行分隔
        
        # 输出统计信息
        print("=" * 50)
        print(f"📊 处理完成！")
        print(f"   ✅ 成功：{success_count} 个")
        if skipped_count > 0:
            print(f"   ⚠️  跳过：{skipped_count} 个（无效文件）")
        if failed_count > 0:
            print(f"   ❌ 失败：{failed_count} 个")


if __name__ == "__main__":
    main()
