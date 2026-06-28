import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.styles.numbers import is_date_format
from datetime import datetime


# ──────────────────────────────────────────────
#  Helpers
# ──────────────────────────────────────────────

def normalize_headers(headers):
    """处理重复 / 空列名，确保 JSON 列名唯一且合法。"""
    seen = {}
    new_headers = []
    for i, h in enumerate(headers):
        h = str(h).strip()
        if not h:
            h = f"列{i + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h}_{seen[h]}"
        else:
            seen[h] = 1
        new_headers.append(h)
    return new_headers


def sanitize_filename(name):
    """过滤文件名中的非法字符（Windows 不允许 \\ / : * ? \" < > |）。"""
    return "".join(c if c not in '\\/:*?"<>|' else "_" for c in name)


def format_cell(cell):
    """将 openpyxl 单元格值按其格式转换为字符串。"""
    if cell.value is None:
        return ""
    fmt = cell.number_format
    val = cell.value
    try:
        if is_date_format(fmt) and isinstance(val, datetime):
            return val.strftime("%Y-%m-%d")
        elif "%" in fmt and isinstance(val, (int, float)):
            return f"{val * 100:.2f}%"
        elif isinstance(val, float):
            if "." in fmt:
                return f"{val:.2f}"
            else:
                return str(round(val, 2))
        else:
            return str(val)
    except Exception as e:
        print(f"Error formatting cell {cell.coordinate}: {e}")
        return str(val)


def unmerge_and_fill_preserve_format(sheet):
    """拆分合并单元格并向下填充值，然后返回格式化后的二维列表。"""
    merged_ranges = list(sheet.merged_cells.ranges)
    for merge_range in merged_ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merge_range))
        value = sheet.cell(row=min_row, column=min_col).value
        sheet.unmerge_cells(str(merge_range))
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                cell = sheet.cell(row=row, column=col)
                if cell.value is None:
                    cell.value = value

    formatted_data = []
    for row in sheet.iter_rows():
        formatted_row = [format_cell(cell) for cell in row]
        formatted_data.append(formatted_row)
    return formatted_data


# ──────────────────────────────────────────────
#  Output writers (shared by xlsx & xls paths)
# ──────────────────────────────────────────────

def _write_outputs(formatted_data, base_name, safe_sheetname, sheet_index,
                   output_dir, exported_files, md_lines):
    """将一个 sheet 的数据写成 CSV / JSON / Markdown 片段。"""

    # --- CSV ---
    output_file = os.path.join(output_dir, f"{base_name}_Sheet{sheet_index}.csv")
    df_csv = pd.DataFrame(formatted_data)
    df_csv.to_csv(output_file, index=False, header=False, encoding="utf-8-sig")
    exported_files.append(output_file)

    # --- JSON (with deduplicated headers) ---
    if formatted_data:
        raw_headers = [str(h) for h in formatted_data[0]]
        headers = normalize_headers(raw_headers)
        if len(formatted_data) > 1:
            json_df = pd.DataFrame(formatted_data[1:], columns=headers)
        else:
            json_df = pd.DataFrame(columns=headers)
    else:
        headers = []
        json_df = pd.DataFrame()

    json_file = os.path.join(output_dir, f"{base_name}_{safe_sheetname}.json")
    json_df.to_json(json_file, orient="records", force_ascii=False)
    exported_files.append(json_file)

    # --- Markdown fragment ---
    if formatted_data and headers:
        md_lines.append("| " + " | ".join(headers) + " |")
        md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
        for row in formatted_data[1:]:
            row_str = [str(cell) for cell in row]
            md_lines.append("| " + " | ".join(row_str) + " |")
    md_lines.append("")


# ──────────────────────────────────────────────
#  Main entry point
# ──────────────────────────────────────────────

def process_excel_file(file_path):
    print(f"Processing Excel file: {file_path}")
    ext = os.path.splitext(file_path)[1].lower()
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    output_dir = os.path.dirname(file_path)
    exported_files = []
    md_lines = []

    # ── .xlsx / .xlsm ─────────────────────────
    if ext in (".xlsx", ".xlsm"):
        try:
            wb = load_workbook(file_path, data_only=True)
        except Exception as e:
            raise Exception(f"Failed to load Excel file: {e}")

        for i, sheetname in enumerate(wb.sheetnames):
            try:
                sheet = wb[sheetname]
                formatted_data = unmerge_and_fill_preserve_format(sheet)

                safe_sheetname = sanitize_filename(sheetname)
                md_lines.append(f"### Sheet{i + 1}: {sheetname}")
                _write_outputs(formatted_data, base_name, safe_sheetname,
                               i + 1, output_dir, exported_files, md_lines)
            except Exception as e:
                print(f"Error processing sheet '{sheetname}': {e}")
                continue

    # ── .xls ──────────────────────────────────
    elif ext == ".xls":
        try:
            # xlrd 2.x 只支持 .xls；通过 pandas engine='xlrd' 读取
            dict_df = pd.read_excel(file_path, sheet_name=None, engine="xlrd")
        except Exception as e:
            raise Exception(
                f"无法读取 .xls 文件（{os.path.basename(file_path)}）: {e}\n"
                "请确认已安装 xlrd>=2.0.1，或将文件另存为 .xlsx 后再试。"
            )

        for i, (sheetname, df) in enumerate(dict_df.items()):
            try:
                # 统一转换为字符串二维列表，与 xlsx 路径保持一致
                headers_row = [str(c) for c in df.columns.tolist()]
                data_rows = df.astype(str).values.tolist()
                # 将 'nan' 替换为空字符串
                data_rows = [
                    ["" if v == "nan" else v for v in row]
                    for row in data_rows
                ]
                formatted_data = [headers_row] + data_rows

                safe_sheetname = sanitize_filename(sheetname)
                md_lines.append(f"### Sheet{i + 1}: {sheetname}")
                md_lines.append("*(注: .xls 文件不支持自动填充合并单元格)*")
                _write_outputs(formatted_data, base_name, safe_sheetname,
                               i + 1, output_dir, exported_files, md_lines)
            except Exception as e:
                print(f"Error processing sheet '{sheetname}': {e}")
                continue

    else:
        raise Exception(f"不支持的文件格式: {ext}（仅支持 .xlsx / .xlsm / .xls）")

    # ── Markdown file ──────────────────────────
    if md_lines:
        md_file = os.path.join(output_dir, f"{base_name}.md")
        try:
            with open(md_file, "w", encoding="utf-8") as f:
                f.write("\n".join(md_lines))
            exported_files.append(md_file)
        except Exception as e:
            print(f"Failed to write markdown file: {e}")

    return exported_files
