"""
Word转Markdown转换器模块
将Word文档转换为Markdown格式，并提取表格为CSV
"""
import os
import csv
import zipfile
from docx import Document as DocxDocument
from docx2python import docx2python


def is_valid_docx(file_path):
    """验证文件是否为有效的 DOCX 文件"""
    try:
        # 检查文件是否存在
        if not os.path.exists(file_path):
            return False
        # 检查文件大小（空文件或过小的文件不是有效的 DOCX）
        if os.path.getsize(file_path) < 100:
            return False
        # 尝试作为 ZIP 文件打开
        with zipfile.ZipFile(file_path, 'r') as zip_ref:
            # 检查是否包含 DOCX 必需的文件
            namelist = zip_ref.namelist()
            return '[Content_Types].xml' in namelist and 'word/document.xml' in namelist
    except (zipfile.BadZipFile, OSError, PermissionError):
        return False


def flatten_docx2_text(nested_list):
    """扁平化嵌套列表，提取文本内容"""
    flat = []
    for item in nested_list:
        if isinstance(item, list):
            flat.extend(flatten_docx2_text(item))
        elif isinstance(item, str) and item.strip():
            flat.append(item.strip())
    return flat


def format_markdown_table(table_rows, title=None):
    """格式化为 Markdown 表格"""
    if not table_rows or len(table_rows) < 2:
        return ""
    header = table_rows[0]
    data_rows = table_rows[1:]

    md_lines = []
    if title:
        md_lines.append(f"### {title}\n\n")
    md_lines.append("| " + " | ".join(header) + " |")
    md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    for row in data_rows:
        md_lines.append("| " + " | ".join(row) + " |")
    md_lines.append("\n")
    return "\n".join(md_lines)


def extract_from_docx(file_path):
    """
    提取 DOCX 文件内容并转换为 Markdown 和 CSV
    
    Args:
        file_path: DOCX文件路径
        
    Returns:
        生成文件路径列表 [md_path, csv1_path, csv2_path, ...]
    """
    # 验证文件
    if not is_valid_docx(file_path):
        raise ValueError(f"文件不是有效的 DOCX 格式: {os.path.basename(file_path)}")
    
    generated_files = []
    
    try:
        md_lines = []

        # 提取正文段落
        try:
            with docx2python(file_path) as doc:
                paragraphs = flatten_docx2_text(doc.body)
                if paragraphs:
                    md_lines.append("## 正文内容\n\n")
                    for para in paragraphs:
                        md_lines.append(para + "\n\n")
        except Exception as e:
            print(f"⚠️  提取正文时出错：{e}")
            md_lines.append("## 正文内容\n\n")
            md_lines.append("（正文提取失败）\n\n")

        # 提取表格内容
        try:
            docx = DocxDocument(file_path)
            for idx, table in enumerate(docx.tables):
                rows = []
                max_cols = 0
                for row in table.rows:
                    cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    max_cols = max(max_cols, len(cells))
                    rows.append(cells)

                # 补齐每行列数
                for row in rows:
                    while len(row) < max_cols:
                        row.append("")

                # 添加 Markdown 表格
                table_title = f"表格 {idx + 1}"
                md_table = format_markdown_table(rows, table_title)
                md_lines.append(md_table)

                # 输出 CSV 文件
                base_name = os.path.splitext(os.path.basename(file_path))[0]
                csv_filename = f"{base_name}_table{idx+1}.csv"
                csv_path = os.path.join(os.path.dirname(file_path), csv_filename)
                with open(csv_path, "w", newline='', encoding="utf-8-sig") as f:
                    writer = csv.writer(f)
                    for row in rows:
                        writer.writerow(row)
                generated_files.append(csv_path)

            # 保存 Markdown 文件
            md_path = os.path.splitext(file_path)[0] + ".md"
            with open(md_path, "w", encoding="utf-8") as md_file:
                md_file.writelines(md_lines)
            
            # Markdown文件放在列表开头
            generated_files.insert(0, md_path)
            
            print(f"✅ 已处理：{os.path.basename(file_path)}")
            print(f"   ➤ Markdown 输出：{md_path}")
            print(f"   ➤ 表格数：{len(docx.tables)}")
            
        except Exception as e:
            print(f"⚠️  提取表格时出错：{e}")
            # 仍然尝试保存已提取的正文内容
            md_path = os.path.splitext(file_path)[0] + ".md"
            with open(md_path, "w", encoding="utf-8") as md_file:
                md_file.writelines(md_lines)
            generated_files.insert(0, md_path)
            print(f"✅ 已保存部分内容：{os.path.basename(file_path)}")

    except Exception as e:
        print(f"❌ 处理文件失败：{os.path.basename(file_path)}")
        print(f"   错误详情：{str(e)}")
        raise
    
    return generated_files
